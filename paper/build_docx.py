#!/usr/bin/env python3
"""Lucrarea in formatul oficial ConsILR-2026 (Word), din ACELASI text ca varianta LaTeX.

Citeste paper.tex (generat de build_paper.py din paper_template.tex) si references.bib, porneste de la
documentul oficial "Instructions for ConsILR-2026.docx" (pagina A4, margini, tema), ii goleste continutul si
scrie lucrarea cu stilurile din Tabelul 1 al instructiunilor (ArticleTitle, Authors, Affiliation, AffiliationEmail,
AbstractTitle, AbstractText, Keywords, Heading 1/2, Normal, NormalFirstParagraph, Caption, ReferencesTitle,
ReferencesItems). Legendele tabelelor stau deasupra, ale figurilor dedesubt ("Figure 1. ..."). Apoi PDF cu LibreOffice.

  cd paper && python3 build_paper.py && python3 build_docx.py
"""
import re, subprocess, sys, unicodedata
import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEMPLATE = "template2026/Instructions for ConsILR-2026.docx"
OUT = "HiddenCore_ConsILR2026"
TNR = "Times New Roman"
TEXT_W = 16.0  # cm, latimea zonei de text (A4 cu margini de 2,5 cm)
AFFIL = ["“Ovidius” University of Constanța, Faculty of Mathematics and Informatics",
         "Bd. Mamaia 124, Constanța, Romania"]
EMAILS = "vlad.belciug@365.univ-ovidius.ro, epelican@univ-ovidius.ro"


# ---------------------------------------------------------------- LaTeX -> segmente de text formatat
def group(s, i, o="{", c="}"):
    """Continutul grupului care incepe la s[i] (acolada sau paranteza dreapta) si pozitia de dupa el."""
    assert s[i] == o, (s[i:i + 40])
    d = 0
    for j in range(i, len(s)):
        if s[j] == o: d += 1
        elif s[j] == c:
            d -= 1
            if d == 0: return s[i + 1:j], j + 1
    raise ValueError("acolada neinchisa: " + s[i:i + 60])


ACC = {("^", "a"): "â", ("^", "A"): "Â", ("^", "i"): "î", ("^", "I"): "Î", ("u", "a"): "ă", ("u", "A"): "Ă",
       ("c", "s"): "ș", ("c", "S"): "Ș", ("c", "t"): "ț", ("c", "T"): "Ț", ("'", "e"): "é", ("'", "E"): "É",
       ("'", "a"): "á", ("'", "i"): "í", ("`", "a"): "à", ("`", "e"): "è", ('"', "a"): "ä", ('"', "e"): "ë",
       ('"', "o"): "ö", ('"', "u"): "ü"}


def accents(s):
    f = lambda m: ACC[(m.group(1), m.group(2).lstrip("\\"))]
    s = re.sub(r"\\([\^'`\"])\{?(\\?[a-zA-Z])\}?", f, s)
    return re.sub(r"\\([uc])\{(\\?[a-zA-Z])\}", f, s)


SYM = {"rightarrow": " → ", "geq": " ≥ ", "leq": " ≤ ", "sum": "∑", "in": " ∈ ", "times": " × ", "approx": " ≈ ",
       "min": "min ", "max": "max "}


def merge(segs):
    out = []
    for t, f in segs:
        if out and out[-1][1] == f: out[-1] = (out[-1][0] + t, f)
        elif t: out.append((t, f))
    return out


def math(s, fmt):
    segs, i = [], 0
    while i < len(s):
        ch = s[i]
        if s.startswith(("\\{", "\\}"), i):
            segs.append((s[i + 1], fmt)); i += 2
        elif ch == "\\":
            m = re.match(r"\\([a-zA-Z]+)", s[i:]); name = m.group(1); i += len(m.group(0))
            if name not in SYM: raise SystemExit(f"simbol matematic necunoscut: \\{name}")
            segs.append((SYM[name], fmt))
        elif ch in "^_":
            if s[i + 1] == "{": inner, i = group(s, i + 1)
            else: inner, i = s[i + 1], i + 2
            segs += math(inner, {**fmt, ("sup" if ch == "^" else "sub"): True})
        elif ch in "<>=":
            segs.append((f" {ch} ", fmt)); i += 1
        elif ch.isalpha():
            segs.append((ch, {**fmt, "i": True})); i += 1
        elif ch == "-":
            segs.append(("−", fmt)); i += 1
        elif ch in "{}":
            i += 1
        else:
            segs.append((ch, fmt)); i += 1
    return [(re.sub(r" {2,}", " ", t), f) for t, f in segs]


class Ctx:
    bib, refs, cited = {}, {}, []


def cite_label(key, paren):
    e = Ctx.bib[key]
    if key not in Ctx.cited: Ctx.cited.append(key)
    fam = [family(a) for a in authors(e["author"])]
    who = fam[0] if len(fam) == 1 else (f"{fam[0]} and {fam[1]}" if len(fam) == 2 and fam[1] != "others"
                                        else f"{fam[0]} et al.")
    y = plain(e["year"])
    return f"{who}, {y}" if paren else f"{who} ({y})"


def inline(s, fmt=None):
    fmt = fmt or {}
    segs, buf, i = [], "", 0
    flush = lambda: segs.append((buf, fmt)) if buf else None
    while i < len(s):
        ch = s[i]
        if ch == "\\":
            m = re.match(r"\\([a-zA-Z]+|.)", s[i:]); name = m.group(1); j = i + len(m.group(0))
            if name in ("ro", "emph", "textit", "textbf", "texttt", "url", "ref", "label", "cite", "newcite", "nocite"):
                opt = None
                if name in ("cite", "newcite") and s[j] == "[": opt, j = group(s, j, "[", "]")
                arg, j = group(s, j)
                flush(); buf = ""
                if name in ("ro", "emph", "textit"):
                    segs += inline(arg, {**fmt, "i": not fmt.get("i"), **({"lang": "ro-RO"} if name == "ro" else {})})
                elif name == "textbf": segs += inline(arg, {**fmt, "b": True})
                elif name == "texttt": segs += inline(arg, {**fmt, "mono": True})
                elif name == "url": segs.append((arg, {**fmt, "url": arg}))
                elif name == "ref": segs.append((str(Ctx.refs[arg]), fmt))
                elif name == "cite":
                    labs = [cite_label(k.strip(), True) for k in arg.split(",")]
                    post = "".join(t for t, _ in inline(opt)) if opt else ""
                    segs.append(("(" + "; ".join(labs) + (f", {post}" if post else "") + ")", fmt))
                elif name == "newcite": segs.append((cite_label(arg.strip(), False), fmt))
                elif name == "nocite": Ctx.cited.extend(k for k in arg.split(",") if k not in Ctx.cited)
                i = j; continue
            if name in ("small", "centering", "normalsize", "hfill"): i = j; continue
            rep = {"%": "%", "&": "&", "_": "_", " ": " ", ",": "\u2009", "break": "\n", "ldots": "…", "-": "",
                   "\\": "\n"}
            if name not in rep: raise SystemExit(f"comanda LaTeX netratata: \\{name} in: {s[max(0, i - 30):i + 30]}")
            buf += rep[name]; i = j
            if name.isalpha() and i < len(s) and s[i] == " " and rep[name] == "": i += 1
            continue
        if ch == "$":
            j = s.index("$", i + 1); flush(); buf = ""
            segs += math(s[i + 1:j], fmt); i = j + 1; continue
        for a, b in (("---", "—"), ("--", "–"), ("``", "“"), ("''", "”"), ("`", "‘"), ("'", "’"), ("~", "\u00a0")):
            if s.startswith(a, i): buf += b; i += len(a); break
        else:
            if ch not in "{}": buf += ch
            i += 1
    flush()
    return merge([(re.sub(r"[ \t\n]+", " ", t) if "\n" not in t.strip(" ") else t, f) for t, f in segs])


def plain(s):
    return "".join(t for t, _ in inline(accents(s)))


# ---------------------------------------------------------------- bibliografia
def parse_bib(path):
    s, out = open(path, encoding="utf8").read(), {}
    for m in re.finditer(r"@(\w+)\{([^,]+),", s):
        body, _ = group(s, s.index("{", m.start()))
        e = {"type": m.group(1).lower()}
        k = 0; body = body[len(m.group(2)) + 1:]
        while True:
            f = re.search(r"(\w+)\s*=\s*\{", body[k:])
            if not f: break
            val, k2 = group(body, k + f.end() - 1)
            e[f.group(1).lower()] = re.sub(r"\s+", " ", val.strip()); k = k2
        out[m.group(2).strip()] = e
    return out


def authors(s):
    return [a.strip() for a in re.split(r"\s+and\s+", s)]


def family(a):
    a = a.strip()
    if a.startswith("{") and a.endswith("}") and a.count("{") == a.count("}"):
        inner, end = group(a, 0)
        if end == len(a): return plain(inner)                       # autor colectiv: {dexonline}
    m = re.search(r"\{([^{}]*)\}$", a)
    return plain(m.group(1)) if m else plain(a.split()[-1])


def names(e):
    A = [plain(a) for a in authors(e["author"])]
    if A[-1] == "others": return ", ".join(A[:-1]) + " et al."
    return A[0] if len(A) == 1 else (" and ".join(A) if len(A) == 2 else ", ".join(A[:-1]) + ", and " + A[-1])


def ref_segments(e):
    """Un element de bibliografie in stilul Computational Linguistics (ca exemplele din sablon)."""
    seg = lambda x, **f: inline(accents(x), f)
    end = lambda t: t if t.rstrip()[-1:] in ".?!" else t + "."
    dot = lambda t: t if t.endswith(".") else t + "."
    S = [(dot(names(e)) + " " + dot(plain(e["year"])) + " ", {})]
    t, ty = e["type"], plain(e["title"])
    if t == "book":
        S += seg(e["title"], i=True) + [(". " if not ty.endswith((".", "?")) else " ", {})]
        S += [(end(", ".join(plain(e[x]) for x in ("publisher", "address") if x in e)) + " ", {})]
    else:
        S += seg(e["title"]) + [(". " if not ty.endswith((".", "?")) else " ", {})]
        if t == "article":
            vol = e.get("volume", "") + (f"({e['number']})" if "number" in e else "")
            S += seg(e["journal"], i=True) + [(", " + vol + (":" + plain(e["pages"]) if "pages" in e else "") + ". ", {})]
        elif t == "inproceedings":
            S += [("In ", {})] + seg(e["booktitle"], i=True)
            tail = ", ".join(x for x in ([f"pages {plain(e['pages'])}"] if "pages" in e else []) +
                             ([plain(e["address"])] if "address" in e else []))
            S += [((", " + tail if tail else "") + ". ", {})]
        elif t == "incollection":
            S += [(f"In {names({'author': e['editor']})} (eds.), ", {})] + seg(e["booktitle"], i=True)
            S += [(f", pages {plain(e['pages'])}. " if "pages" in e else ". ", {})]
            S += [(end(", ".join(plain(e[x]) for x in ("publisher", "address") if x in e)) + " ", {})]
        elif t == "misc":
            S += seg(e["howpublished"]) + [(". ", {})]
    if "note" in e: S += seg(e["note"])
    S = merge(S)
    S[-1] = (S[-1][0].rstrip(), S[-1][1])
    return S


def sort_key(k):
    e = Ctx.bib[k]
    f = unicodedata.normalize("NFKD", family(authors(e["author"])[0])).encode("ascii", "ignore").decode().lower()
    return (f, plain(e["year"]))


# ---------------------------------------------------------------- stiluri (Tabelul 1 din instructiuni)
def set_rfonts(rpr, name):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in list(rf.attrib): del rf.attrib[a]
    for a in ("ascii", "hAnsi", "eastAsia", "cs"): rf.set(qn(f"w:{a}"), name)


def style(doc, name, size, bold=False, italic=False, caps=False, font=TNR, align=None, first=None, left=None,
          right=None, hanging=None, before=0, after=0, keep=False, outline=None, base="Normal", line=True):
    st = doc.styles[name] if name in [s.name for s in doc.styles] else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if name != "Normal": st.base_style = doc.styles[base]
    st.hidden = False; st.quick_style = True
    f = st.font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.all_caps = caps; f.color.rgb = RGBColor(0, 0, 0)
    set_rfonts(st.element.get_or_add_rPr(), font)
    p = st.paragraph_format
    p.alignment = align; p.space_before = Pt(before); p.space_after = Pt(after); p.keep_with_next = keep
    p.left_indent = Cm(left) if left is not None else Cm(0); p.right_indent = Cm(right) if right is not None else Cm(0)
    p.first_line_indent = Cm(-hanging) if hanging else (Cm(first) if first is not None else Cm(0))
    if line: p.line_spacing = 1.05
    ppr = st.element.get_or_add_pPr()
    for tag in ("w:numPr", "w:pBdr"):
        for el in ppr.findall(qn(tag)): ppr.remove(el)
    for el in ppr.findall(qn("w:outlineLvl")): ppr.remove(el)
    if outline is not None:
        o = OxmlElement("w:outlineLvl"); o.set(qn("w:val"), str(outline)); ppr.append(o)
    return st


def setup_styles(doc):
    rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    set_rfonts(rpr, TNR)
    for el in rpr.findall(qn("w:sz")) + rpr.findall(qn("w:szCs")): el.set(qn("w:val"), "22")
    J, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER
    style(doc, "Normal", 11, align=J, first=0.4)
    style(doc, "NormalFirstParagraph", 11, align=J, first=0)
    style(doc, "ArticleTitle", 15, bold=True, caps=True, align=C, after=0)
    style(doc, "Authors", 12, caps=True, align=C)
    style(doc, "Affiliation", 11, italic=True, align=C)
    style(doc, "AffiliationEmail", 11, font="Courier New", align=C)
    style(doc, "AbstractTitle", 12, bold=True, align=C, before=18, after=12, keep=True)
    style(doc, "AbstractText", 11, align=J, left=0.6, right=0.6, after=6)
    style(doc, "Keywords", 10, align=J, left=0.6, right=0.6, before=6, after=6)
    style(doc, "Heading 1", 12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=12, keep=True, outline=0)
    style(doc, "Heading 2", 11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=8, keep=True, outline=1)
    style(doc, "Caption", 11, align=C, before=6, after=6)
    style(doc, "AcknowledgementsTitle", 11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, keep=True)
    style(doc, "ReferencesTitle", 11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, keep=True)
    style(doc, "ReferencesItems", 10, align=WD_ALIGN_PARAGRAPH.LEFT, hanging=0.4, after=6)   # URL-urile lungi nu mai lasa goluri
    style(doc, "TableText", 10, align=WD_ALIGN_PARAGRAPH.LEFT, line=False)
    style(doc, "FigureParagraph", 11, align=C, before=6, after=0, keep=True, line=False)


def clear_body(doc):
    body = doc.element.body
    for ch in list(body):
        if ch.tag != qn("w:sectPr"): body.remove(ch)
    used = set(re.findall(r'r:(?:embed|id)="([^"]+)"', doc.element.xml))
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype in (RT.IMAGE, RT.HYPERLINK) and rid not in used:
            del doc.part.rels[rid]
    # notele de subsol ale sablonului (raman doar separatorii)
    for rel in doc.part.rels.values():
        if rel.reltype == RT.FOOTNOTES:
            part = rel.target_part
            xml = part.blob.decode("utf8")
            xml = re.sub(r'<w:footnote (?![^>]*w:type="(?:separator|continuationSeparator)")[^>]*>.*?</w:footnote>', "",
                         xml, flags=re.S)
            part._blob = xml.encode("utf8")


# ---------------------------------------------------------------- scrierea textului
def add_hyperlink(par, text, url, size=None):
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "000000"); rpr.append(col)
    # spatiu de latime zero dupa / si _ : URL-urile lungi se pot rupe la capat de rand fara goluri in text
    shown = re.sub(r"([/_])(?=.)", "\\1\u200b", text.replace("https://", "\0")).replace("\0", "https://")
    r.append(rpr); t = OxmlElement("w:t"); t.text = shown; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); par._p.append(h)


def write(par, segs, url_size=None):
    for text, f in segs:
        if f.get("url"):
            add_hyperlink(par, text, f["url"], url_size); continue
        parts = text.split("\n")
        for k, piece in enumerate(parts):
            if k: par.add_run().add_break(WD_BREAK.LINE)
            if not piece: continue
            r = par.add_run(piece)
            if f.get("i"): r.italic = True
            if f.get("b"): r.bold = True
            if f.get("sup"): r.font.superscript = True
            if f.get("sub"): r.font.subscript = True
            if f.get("mono"): r.font.name = "Courier New"
            if f.get("lang"):
                lang = OxmlElement("w:lang"); lang.set(qn("w:val"), f["lang"]); r._r.get_or_add_rPr().append(lang)
    return par


def para(doc, style_name, segs, **kw):
    return write(doc.add_paragraph(style=style_name), segs, **kw)


def cell_border(cell, side, sz):
    tcpr = cell._tc.get_or_add_tcPr()
    b = tcpr.find(qn("w:tcBorders"))
    if b is None: b = OxmlElement("w:tcBorders"); tcpr.append(b)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
    b.append(el)


def no_table_borders(table):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:val"), "nil"); b.append(el)
    tblpr.append(b)


TNR_TTF = "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf"


def text_cm(t, size):
    """Latimea textului in Times New Roman (masurata cu metricile fontului; aproximare daca fontul lipseste)."""
    try:
        from PIL import ImageFont
        f = ImageFont.truetype(TNR_TTF, size * 20)
        return f.getlength(t) / 20 * 2.54 / 72
    except Exception:
        return len(t) * 0.19 * size / 10


def add_table(doc, spec, rows, size):
    """Tabel in stil booktabs: linie groasa sus si jos, linie subtire sub antet; | din specificatie -> linie verticala."""
    cols, bars = [], set()
    for ch in spec:
        if ch == "|": bars.add(len(cols))
        elif ch in "lrc": cols.append(ch)
    data = [r for r in rows if r != "MIDRULE"]
    mids = [k - 1 - rows[:k].count("MIDRULE") for k, r in enumerate(rows) if r == "MIDRULE"]
    t = doc.add_table(rows=len(data), cols=len(cols)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_table_borders(t)
    cells_txt = [["".join(x for x, _ in c) for c in r] for r in data]
    w = [max(text_cm(r[j], size) for r in cells_txt if j < len(r)) * 1.05 + 0.6 for j in range(len(cols))]
    while sum(w) > TEXT_W:                  # se ingusteaza doar coloana cea mai lata (textul ei se rupe pe doua randuri)
        j = max(range(len(w)), key=lambda x: w[x]); w[j] -= 0.05
    scale = 1.0
    tblpr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblpr.append(lay)
    AL = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}
    for i, r in enumerate(data):
        for j in range(len(cols)):
            c = t.cell(i, j); c.width = Cm(w[j] * scale)
            p = c.paragraphs[0]; p.style = doc.styles["TableText"]; p.alignment = AL[cols[j]]
            for run in write(p, r[j] if j < len(r) else []).runs: run.font.size = Pt(size)
            if i == 0: cell_border(c, "top", 12)
            if i == len(data) - 1: cell_border(c, "bottom", 12)
            if i in mids: cell_border(c, "bottom", 4)
            if j in bars: cell_border(c, "left", 4)
            # tabelele scurte raman pe o pagina; cele lungi se pot continua, cu antetul repetat
            if i < len(data) - 1 and (len(data) <= 12 or i == 0): p.paragraph_format.keep_with_next = True
        trpr = t.rows[i]._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trpr.append(cs)
        if i == 0 and len(data) > 12: trpr.append(OxmlElement("w:tblHeader"))
    for j, col in enumerate(t.columns): col.width = Cm(w[j] * scale)
    return t


def caption(doc, kind, n, segs):
    return para(doc, "Caption", [(f"{kind} {n}. ", {})] + segs)


def render_float(doc, env, n):
    cap, _ = group(env, env.index("\\caption") + len("\\caption"))
    if env.startswith("\\begin{table}"):
        size = 9 if "\\small" in env else 10
        m = re.search(r"\\begin\{tabular\}\{([^}]*)\}(.*?)\\end\{tabular\}", env, re.S)
        rows = []
        for raw in re.split(r"\\\\", m.group(2)):
            for line in raw.strip().split("\n"):
                line = line.strip()
                if not line: continue
                if line in ("\\toprule", "\\bottomrule"): continue
                if line == "\\midrule": rows.append("MIDRULE"); continue
                if line.startswith("\\midrule"): rows.append("MIDRULE"); line = line[len("\\midrule"):].strip()
                if line: rows.append([inline(c.strip()) for c in re.split(r"(?<!\\)&", line)])
        cp = caption(doc, "Table", n, inline(cap)); cp.paragraph_format.keep_with_next = True
        add_table(doc, m.group(1), rows, size)
        doc.add_paragraph(style="TableText").paragraph_format.space_after = Pt(4)
    else:
        m = re.search(r"\\includegraphics\[width=([\d.]+)\\textwidth\]\{([^}]+)\}", env)
        p = doc.add_paragraph(style="FigureParagraph")
        p.add_run().add_picture("../" + m.group(2).replace(".pdf", ".png"), width=Cm(float(m.group(1)) * TEXT_W))
        caption(doc, "Figure", n, inline(cap))


# ---------------------------------------------------------------- documentul
def main():
    tex = open("paper.tex", encoding="utf8").read()
    Ctx.bib = parse_bib("references.bib")
    envs = re.findall(r"\\begin\{(table|figure)\}.*?\\end\{\1\}", tex, re.S)
    nt = nf = 0
    for m in re.finditer(r"\\begin\{(table|figure)\}(.*?)\\end\{\1\}", tex, re.S):
        if m.group(1) == "table": nt += 1; n = nt
        else: nf += 1; n = nf
        for lab in re.findall(r"\\label\{([^}]+)\}", m.group(2)): Ctx.refs[lab] = n
    doc = docx.Document(TEMPLATE)
    clear_body(doc); setup_styles(doc)
    title, _ = group(tex, tex.index("\\title{") + 6)
    para(doc, "ArticleTitle", inline(title))
    doc.add_paragraph(style="Authors")
    auths = [plain(a) for a in re.findall(r"\\author\[\d\]\{([^}]*)\}", tex)]
    para(doc, "Authors", [(", ".join(auths), {})])
    doc.add_paragraph(style="Authors")
    for line in AFFIL: para(doc, "Affiliation", [(line, {})])
    para(doc, "AffiliationEmail", [(EMAILS, {})])
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", tex, re.S).group(1).strip()
    para(doc, "AbstractTitle", [("Abstract", {})])
    ab = inline(abstract); para(doc, "AbstractText", ab)
    nw = len("".join(t for t, _ in ab).split())
    kw = re.search(r"\\keyword\{\}(.*?)\n\s*\n", tex, re.S).group(1).strip()
    para(doc, "Keywords", [("Keywords: ", {"b": True})] + inline(kw))

    body = tex[tex.index("\\section{"):tex.index("\\bibliographystyle")]
    floats = {}
    def stash(m):
        key = f"@@FLOAT{len(floats)}@@"; floats[key] = m.group(0); return key
    body = re.sub(r"\\begin\{(table|figure)\}.*?\\end\{\1\}", stash, body, flags=re.S)
    counters = {"table": 0, "figure": 0}
    def emit_float(key):
        env = floats[key]; kind = "table" if env.startswith("\\begin{table}") else "figure"
        counters[kind] += 1; render_float(doc, env, counters[kind])
    sec = sub = 0; first = True
    for block in re.split(r"\n\s*\n", body):
        block = block.strip()
        if not block: continue
        while True:
            m = re.match(r"\\(section\*?|subsection)\{", block)
            if not m: break
            head, j = group(block, m.end() - 1); kind = m.group(1)
            if kind == "section": sec += 1; sub = 0; para(doc, "Heading 1", [(f"{sec}. ", {})] + inline(head))
            elif kind == "section*": para(doc, "AcknowledgementsTitle", inline(head))
            else: sub += 1; para(doc, "Heading 2", [(f"{sec}.{sub}. ", {})] + inline(head))
            first = True; block = block[j:].strip()
        if not block: continue
        keys = re.findall(r"@@FLOAT\d+@@", block)
        text = re.sub(r"\s*@@FLOAT\d+@@\s*", " ", block).strip()
        if text:
            para(doc, "NormalFirstParagraph" if first else "Normal", inline(text)); first = False
        for k in keys: emit_float(k)
    para(doc, "ReferencesTitle", [("References", {})])
    for k in sorted(dict.fromkeys(Ctx.cited), key=sort_key):
        para(doc, "ReferencesItems", ref_segments(Ctx.bib[k]), url_size=10)
    doc.core_properties.title = plain(title).title()
    doc.core_properties.author = ", ".join(a.title() for a in auths)
    doc.core_properties.language = "en-GB"
    doc.save(OUT + ".docx")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", OUT + ".docx"], capture_output=True)
    info = subprocess.run(["pdfinfo", OUT + ".pdf"], capture_output=True, text=True).stdout
    pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    print(f"{OUT}.docx/.pdf: {pages} pagini ({'par' if pages % 2 == 0 else 'IMPAR - trebuie numar par'}) | "
          f"abstract {nw} cuvinte (max 200) | tabele {nt}, figuri {nf} | referinte {len(set(Ctx.cited))}")
    if nw > 200: sys.exit("abstractul depaseste 200 de cuvinte")


if __name__ == "__main__":
    main()
