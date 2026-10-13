#!/usr/bin/env python3
"""Prezentarea ConsILR 2026 (pptx) din results/results.json. Note de vorbitor in romana.

  cd slides && python3 make_slides.py && soffice --headless --convert-to pdf HiddenCore_ConsILR2026.pptx
"""
import json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

R = json.load(open("../results/results.json", encoding="utf8"))
N = R["data"]["headwords"]; ai = R["regimes"]["ai"]
LB = ai.get("minset_lb_ilp", ai["minset_lb"]); UB = ai["minset_ub"]
_rg = R["regimes"].values()
RNG_K = f"{min(100 * r['kernel'] / N for r in _rg):.0f}–{max(100 * r['kernel'] / N for r in _rg):.0f}%"
RNG_M = f"{min(100 * r.get('minset_lb_ilp', r['minset_lb']) / N for r in _rg):.1f}–{max(100 * r['minset_ub'] / N for r in _rg):.1f}%"
G, S, E = R["graur"], R["swadesh"], R["etymology"]
LOGO = "../../../ConsILR2025/logo_consilr.png"
FIG = "../figures/"
INK, INK2, MUTED = RGBColor(0x0b, 0x0b, 0x0b), RGBColor(0x52, 0x51, 0x4e), RGBColor(0x8a, 0x89, 0x84)
BLUE, ORANGE, AQUA = RGBColor(0x2a, 0x78, 0xd6), RGBColor(0xeb, 0x68, 0x34), RGBColor(0x1b, 0xaf, 0x7a)
SURF, LINE = RGBColor(0xfc, 0xfc, 0xfb), RGBColor(0xe6, 0xe5, 0xe1)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def p(n, d=1):
    return f"{100 * n:.{d}f}%"


def text(slide, x, y, w, h, runs, size=22, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        segs = para if isinstance(para, tuple) else (para,)
        for seg in segs:
            r = par.add_run()
            if isinstance(seg, dict):
                r.text = seg["t"]; f = r.font
                f.size = Pt(seg.get("size", size)); f.bold = seg.get("bold", bold); f.italic = seg.get("italic", False)
                f.color.rgb = seg.get("color", color); f.name = FONT
            else:
                r.text = seg; f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = FONT
        par.space_after = Pt(size * 0.45)
    return tb


def base(title=None, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = SURF
    s.shapes.add_picture(LOGO, W - Inches(1.75), Inches(0.28), height=Inches(0.55))
    if kicker:
        text(s, Inches(0.7), Inches(0.35), Inches(10), Inches(0.4), kicker.upper(), size=13, color=BLUE, bold=True)
    if title:
        text(s, Inches(0.7), Inches(0.68), Inches(11.2), Inches(1.0), title, size=32, bold=True)
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.7), H - Inches(0.45), W - Inches(0.7), H - Inches(0.45))
    ln.line.color.rgb = LINE; ln.line.width = Pt(0.75)
    text(s, Inches(0.7), H - Inches(0.42), Inches(9), Inches(0.35),
         "Belciug & Pelican · The Hidden Core of Romanian · ConsILR 2026", size=11, color=MUTED)
    return s


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def box(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    b = s.shapes.add_shape(shape, x, y, w, h)
    b.fill.solid(); b.fill.fore_color.rgb = fill
    if line: b.line.color.rgb = line
    else: b.line.fill.background()
    b.shadow.inherit = False
    return b


def stat(s, x, y, big, small, color=BLUE, w=Inches(3.6)):
    text(s, x, y, w, Inches(1.0), big, size=54, bold=True, color=color)
    text(s, x, y + Inches(1.0), w, Inches(1.0), small, size=17, color=INK2)


# 1 ── titlu
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = SURF
s.shapes.add_picture(LOGO, Inches(0.7), Inches(0.6), height=Inches(0.9))
text(s, Inches(0.7), Inches(2.0), Inches(11.5), Inches(1.6), "The Hidden Core of Romanian", size=54, bold=True)
text(s, Inches(0.7), Inches(3.15), Inches(11.5), Inches(1.0), "Latent structure of the DEX and Graur’s fundamental lexical stock",
     size=28, color=INK2)
text(s, Inches(0.7), Inches(4.6), Inches(11.5), Inches(1.4),
     [({"t": "Vlad Belciug, Elena Pelican", "bold": True, "size": 22},),
      "“Ovidius” University of Constanța · Faculty of Mathematics and Informatics",
      "ConsILR 2026 · Galați & online · 9–11 December 2026"], size=18, color=INK2)
notes(s, "Bună ziua. Vă propun o întrebare simplă despre dicționar și un răspuns măsurat pe tot DEX-ul.")

# 2 ── cercul
s = base("A dictionary is a circle", "The question")
cop = box(s, Inches(1.4), Inches(2.6), Inches(3.2), Inches(1.3), RGBColor(0xe8, 0xf1, 0xfb), BLUE)
arb = box(s, Inches(7.2), Inches(2.6), Inches(3.2), Inches(1.3), RGBColor(0xe8, 0xf1, 0xfb), BLUE)
for b, w in ((cop, "copac"), (arb, "arbore")):
    tf = b.text_frame; tf.text = w; r = tf.paragraphs[0].runs[0]
    r.font.size = Pt(34); r.font.italic = True; r.font.color.rgb = INK; r.font.name = FONT
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
for (x1, y1, x2, y2, lab, ly) in ((Inches(4.6), Inches(2.95), Inches(7.2), Inches(2.95), "defined as “arbore”", Inches(2.25)),
                                  (Inches(7.2), Inches(3.55), Inches(4.6), Inches(3.55), "defined as “copac”", Inches(3.7))):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = ORANGE; c.line.width = Pt(3)
    c.line._get_or_add_ln().append(c.line._get_or_add_ln().makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"}))
    text(s, Inches(4.7), ly, Inches(2.4), Inches(0.5), lab, size=15, color=INK2, align=PP_ALIGN.CENTER)
text(s, Inches(1.4), Inches(4.6), Inches(10.5), Inches(1.8),
     [({"t": "How many words must you already know ", "bold": True},
       {"t": "before the dictionary can teach you all the others?", "bold": True, "color": BLUE}),
      "Somewhere the circle must be broken by words learned from experience, not from definitions (the symbol grounding problem, Harnad 1990)."],
     size=24)
notes(s, "Cauți copac în DEX: scrie arbore. Cauți arbore: scrie copac. Dacă nu știi niciunul, rămâi în cerc. "
         "Întrebarea: câte cuvinte trebuie să știi dinainte, din experiență, ca să înțelegi tot restul doar din definiții?")

# 3 ── metoda
s = base("Dictionaries as graphs: Kernel, Core, MinSet", "The method (Vincent-Lamarre, Harnad et al. 2016)")
cx, cy = Inches(3.4), Inches(4.2)
for rad, col, lab in ((2.45, RGBColor(0xf0, 0xef, 0xec), "Rest"), (1.75, RGBColor(0xcd, 0xe2, 0xfb), "Kernel"),
                      (1.15, RGBColor(0x86, 0xb6, 0xef), "Core"), (0.55, BLUE, "MinSet")):
    o = box(s, cx - Inches(rad), cy - Inches(rad), Inches(2 * rad), Inches(2 * rad), col, None, MSO_SHAPE.OVAL)
for lab, yy, col in (("Rest", 1.95, INK2), ("Kernel", 2.65, INK2), ("Core", 3.25, INK), ("MinSet", 3.95, SURF)):
    text(s, cx - Inches(1.0), Inches(yy), Inches(2.0), Inches(0.4), lab, size=16, bold=True, color=col, align=PP_ALIGN.CENTER)
text(s, Inches(6.6), Inches(1.9), Inches(6.0), Inches(4.6),
     [({"t": "Graph: ", "bold": True}, "an arrow from every defining word to the word it defines."),
      ({"t": "Kernel: ", "bold": True}, "what remains after repeatedly removing words that define nothing. Everything else can be defined from it."),
      ({"t": "Core: ", "bold": True}, "its largest strongly connected part."),
      ({"t": "MinSet: ", "bold": True}, "the smallest set of words from which the whole dictionary can be learned (an NP-hard problem)."),
      ({"t": "English dictionaries: ", "bold": True, "color": BLUE}, {"t": "Kernel 7–12%, MinSet ≈1%.", "color": BLUE})],
     size=19)
notes(s, "Metoda lui Harnad: dicționarul devine un graf. Nucleul (Kernel) e ce rămâne după ce scoți repetat cuvintele care nu definesc nimic. "
         "MinSet e cel mai mic set de cuvinte din care poți învăța tot dicționarul. În engleză: nucleu 7–12%, MinSet circa 1%.")

# 4 ── Graur
s = base("Romanian linguistics asked a similar question in 1954", "Graur’s fundamental lexical stock")
text(s, Inches(0.7), Inches(1.9), Inches(7.0), Inches(4.5),
     [({"t": "Al. Graur, ", "bold": True}, {"t": "Încercare asupra fondului principal lexical al limbii romîne", "italic": True}, " (1954)"),
      "1,419 words, in three circles of certainty (964 · 217 · 238), each labelled with its origin.",
      "Criteria: used by the whole people, old and stable, productive, rich in expressions.",
      ({"t": "His own test: ", "bold": True}, "a core word occupies a large space in dictionaries (pp. 24–25)."),
      ({"t": "He called the selection subjective ", "color": INK2}, {"t": "(1957). We treat it as a historical hypothesis, not a gold standard.", "color": INK2})],
     size=20)
box(s, Inches(8.3), Inches(2.0), Inches(4.3), Inches(3.9), RGBColor(0xf0, 0xef, 0xec))
text(s, Inches(8.55), Inches(2.2), Inches(3.9), Inches(3.6),
     [({"t": "We transcribed the list", "bold": True},),
      f"{G['entries']:,} entries · Slavic share 21.49% (exactly as published)",
      f"{G['linked']:,} linked to DEX headwords",
      "Graur left out arbore because it competes with copac, pom, lemn (p. 19)."], size=17, color=INK)
notes(s, "Graur a pus în 1954 aceeași întrebare, cu alte mijloace: 1.419 cuvinte, în trei cercuri de siguranță. "
         "El însuși spunea că importanța unui cuvânt se vede în spațiul ocupat în dicționare. Am transcris lista și am reprodus exact totalul și procentul slav de 21,49%.")

# 5 ── date si AI
s = base("Reading the whole DEX, with a language model as an instrument", "Data and method")
stat(s, Inches(0.7), Inches(1.9), f"{R['data']['definitions']:,}", "definitions, DEX 2nd edition (dexonline, GPL)")
stat(s, Inches(4.7), Inches(1.9), f"{N:,}", "headwords (DEX ’98 completed with DEX ’96 for G–I)")
stat(s, Inches(8.7), Inches(1.9), p(ai["ambig_rate"], 0), "of word forms in definitions are ambiguous", color=ORANGE)
text(s, Inches(0.7), Inches(4.25), Inches(12.0), Inches(2.4),
     [({"t": "are", "italic": True}, " → avea or ara?   ", {"t": "fire", "italic": True}, " → fi, fir or fire?"),
      ({"t": "Claude Haiku 4.5 ", "bold": True}, "reads each definition and picks the lemma; a stronger model (Claude Opus 5) checks a random sample: ",
       {"t": f"{p(R['validation']['opus_random_content'], 0)} agreement on content words.", "bold": True, "color": BLUE}),
      f"Cost ≈ {R['validation']['cost_usd']:.0f} USD. Every result is repeated without the model (frequency rule, all candidates)."],
     size=20)
notes(s, "Un sfert din formele din definiții sunt ambigue. Modelul AI citește fiecare definiție și alege lema. "
         "Îl verificăm cu un model mai puternic: 87% acord pe cuvintele de conținut. Totul a costat cât un prânz. Și refacem totul fără AI, ca să vedem dacă concluziile rezistă.")

# 6 ── rezultatul principal
s = base("The DEX has a very large hidden core", "Result 1")
s.shapes.add_picture(FIG + "fig1_size.png", Inches(0.7), Inches(1.9), width=Inches(7.3))
stat(s, Inches(8.6), Inches(1.9), p(ai["kernel"] / N, 0), f"of headwords in the Kernel ({ai['kernel']:,} words)")
stat(s, Inches(8.6), Inches(4.1), f"{100*LB/N:.1f}–{100*UB/N:.1f}%", f"MinSet: {LB:,}–{UB:,} words (proven lower and upper bounds)", color=ORANGE, w=Inches(4.4))
notes(s, f"Nucleul DEX are {p(ai['kernel']/N,0)} din cuvinte, de trei până la cinci ori mai mult decât în engleză. "
         f"Setul minim: între {LB:,} și {UB:,} cuvinte, adică aproape zece procente, de zece ori mai mult decât în engleză. Intervalul e calculat exact, cu programare liniară în numere întregi.")

# 7 ── de ce: perechile
s = base("Why? Definition by synonym", "Result 2")
pairs = R["pairs"]["examples"][:6]
for i, (a, b) in enumerate(pairs):
    x = Inches(0.7 + (i % 3) * 4.1); y = Inches(1.95 + (i // 3) * 1.25)
    bx = box(s, x, y, Inches(3.8), Inches(1.0), RGBColor(0xe8, 0xf1, 0xfb))
    tf = bx.text_frame; tf.text = f"{a}  ⇄  {b}"
    r = tf.paragraphs[0].runs[0]; r.font.size = Pt(24); r.font.italic = True; r.font.color.rgb = INK; r.font.name = FONT
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
text(s, Inches(0.7), Inches(4.65), Inches(12.0), Inches(2.0),
     [({"t": f"{ai['mutual_pairs']:,} pairs of words define each other", "bold": True, "color": BLUE},
       f" ({p(R['pairs']['derivational'],0)} share a stem; most others are synonyms)."),
      f"Each pair forces one word into the MinSet: {p(R.get('two_cycle_share', 0.98), 0)} of the disjoint cycles are such pairs.",
      "English learner dictionaries use a controlled defining vocabulary; the DEX defines by synonym."], size=20)
notes(s, "De ce? Pentru că DEX-ul definește prin sinonime. Aproape nouă mii de perechi (8.761) se definesc reciproc. Fiecare pereche te obligă să știi dinainte unul dintre cuvinte. "
         "Nu sunt greșeli: e un stil de definire. Graur însuși observase concurența arbore–copac.")

# 8 ── robustete
s = base("The conclusion does not depend on the AI", "Robustness")
s.shapes.add_picture(FIG + "fig3_robustness.png", Inches(0.7), Inches(1.9), width=Inches(7.6))
text(s, Inches(8.7), Inches(2.0), Inches(4.2), Inches(4.4),
     ["Three ways of resolving ambiguous forms: the most frequent lemma, the language model, all candidates.",
      ({"t": "Same picture every time: ", "bold": True}, f"a Kernel of {RNG_K} and a MinSet of {RNG_M}."),
      "Correct lemmas reveal more reciprocal pairs."], size=19)
notes(s, f"Refacem tot fără AI, în două feluri. Rezultatul rămâne același: nucleu {RNG_K}, MinSet {RNG_M}. Deci concluzia nu e un artefact al modelului.")

# 9 ── validare cu liste de experti
s = base("The core is linguistically real", "Result 3")
s.shapes.add_picture(FIG + "fig2_enrichment.png", Inches(0.6), Inches(1.85), width=Inches(8.3))
text(s, Inches(9.2), Inches(2.0), Inches(3.8), Inches(4.4),
     [({"t": f"{p(S['layers']['Kernel']['rate'],0)}", "bold": True, "color": BLUE, "size": 30}, " of the Swadesh list (about 200 basic meanings used in comparative linguistics) in the Kernel"),
      ({"t": f"{p(G['all']['Kernel']['rate'],0)}", "bold": True, "color": BLUE, "size": 30}, " of Graur’s words in the Kernel"),
      ({"t": f"chance: {p(R['layers']['Kernel']['share'],0)}", "color": INK2},),
      f"Graur’s most certain words are the most central: {p(G['core']['Kernel']['rate'],0)} in the Kernel, {p(G['core']['MinSet']['rate'],0)} in the MinSet (wider groups: {p(G['middle']['MinSet']['rate'],0)}, {p(G['maximal']['MinSet']['rate'],0)})."],
     size=18)
notes(s, f"Este nucleul real lingvistic? {p(S['layers']['Kernel']['rate'],0)} din lista Swadesh și {p(G['all']['Kernel']['rate'],0)} din lista lui Graur sunt în nucleu, față de {p(R['layers']['Kernel']['share'],0)} la întâmplare. "
         "Lista Swadesh este lista standard de circa 200 de sensuri de bază folosită în lingvistica comparată. Cuvintele pe care Graur le considera cele mai sigure sunt cele mai centrale: 97% în nucleu și 49% în setul minim; cele două cercuri mai largi diferă puțin între ele.")

# 9b ── exemplele lui Graur
WD = json.load(open("html_data.json", encoding="utf8"))["words"]
s = base("Graur’s own examples, seen through the graph", "Case studies")
CATN = {"C": "most certain", "M": "middle circle", "X": "maximal circle", "": "not in the list"}
LAYN = {"minset": "MinSet", "core": "Core", "satellite": "Satellite", "rest": "outside Kernel"}
ex = [w for w in ["arbore", "copac", "pom", "femeie", "nevastă", "muiere", "cioban", "păcurar", "doagă", "săgeată", "concert", "axiomă"] if w in WD]
rows_, cols_ = len(ex) + 1, 4
tbl = s.shapes.add_table(rows_, cols_, Inches(0.7), Inches(1.85), Inches(8.2), Inches(0.36 * rows_)).table
for j, h in enumerate(["word", "Graur 1954", "graph layer", "used in definitions of"]):
    c = tbl.cell(0, j); c.text = h
for i, w in enumerate(ex, 1):
    d = WD[w]
    for j, v in enumerate([w, CATN[d["g"]], LAYN[d["l"]], f"{d['o']:,} headwords"]):
        tbl.cell(i, j).text = v
for i in range(rows_):
    for j in range(cols_):
        c = tbl.cell(i, j); para = c.text_frame.paragraphs[0]
        for r in para.runs:
            r.font.size = Pt(15); r.font.name = FONT; r.font.bold = (i == 0); r.font.italic = (j == 0 and i > 0)
            r.font.color.rgb = SURF if i == 0 else INK
        c.fill.solid(); c.fill.fore_color.rgb = BLUE if i == 0 else (SURF if i % 2 else RGBColor(0xf3, 0xf2, 0xef))
text(s, Inches(9.3), Inches(1.95), Inches(3.7), Inches(4.6),
     [({"t": "arbore", "italic": True, "bold": True}, " was excluded by Graur because it competes with copac, pom, lemn (p. 19)."),
      ({"t": f"Yet the DEX uses arbore in {WD['arbore']['o']} definitions, copac in {WD['copac']['o']}.", "color": BLUE},),
      "The dictionary’s defining preferences and the speakers’ usage Graur described are different things."],
     size=17)
notes(s, "Exemplele lui Graur. A exclus arbore pentru că e în concurență cu copac; dar DEX-ul folosește arbore aproape de două ori mai des decât copac ca termen definitoriu. "
         "Preferințele dicționarului nu coincid cu uzul descris de Graur, și asta e o observație interesantă despre metalimbajul lexicografic.")

# 10 ── doua nuclee
s = base("Two cores: inherited words and working words", "Result 4")
s.shapes.add_picture(FIG + "fig4_etymology.png", Inches(0.6), Inches(1.85), width=Inches(8.2))
outk = [w for w in ["arici", "mânz", "țap", "țarc", "teasc", "topor", "prag", "ginere", "socru", "vișin"] if w in G["outside_kernel"]]
text(s, Inches(9.1), Inches(1.95), Inches(3.9), Inches(4.6),
     [({"t": "Graur’s core: ", "bold": True}, "stable, inherited (Latin + Slavic)."),
      ({"t": "The DEX’s working core: ", "bold": True}, "learned words: acțiune, rezultat, persoană, substanță…"),
      f"Only {E['graur_in_top1419']} of Graur’s words are among the 1,419 most-used defining words.",
      ({"t": "Outside the Kernel: ", "bold": True}, {"t": ", ".join(outk[:7]) + "…", "italic": True},
       " (rural words that define nothing else)")], size=17)
notes(s, "Dar nu e identitate. Cuvintele cu care DEX-ul își scrie definițiile sunt în mare parte împrumuturi savante franceze. "
         "Nucleul lui Graur e unul al stabilității; al dicționarului, unul al muncii de definire. Cuvintele rurale ale lui Graur sunt definite, dar nu definesc nimic.")

# 11 ── ce iau lingvistii
s = base("Implications for Romanian lexicography and linguistics", "Discussion")
items = [("A map of circularity", "the list of mutually defining pairs shows where the DEX relies on synonymy; useful for a revised or learner’s dictionary."),
         ("A test of a classical idea", "Graur’s intuition is confirmed with data he never had, and placed in context."),
         ("AI as an instrument", "a narrow, checkable task at full scale; reliability measured; conclusions tested without the model.")]
for i, (h, b) in enumerate(items):
    y = Inches(1.95 + i * 1.5)
    box(s, Inches(0.7), y, Inches(0.18), Inches(1.15), BLUE)
    text(s, Inches(1.1), y - Inches(0.05), Inches(11.4), Inches(1.3), [({"t": h, "bold": True, "size": 24},), b], size=19)
notes(s, "Trei lucruri pentru lingviști: o hartă a circularității pentru lexicografi, un test cantitativ al unei idei clasice și o rețetă de folosire a AI-ului ca instrument verificabil.")

# 12 ── cat de siguri suntem
s = base("Strength of the evidence", "How confident are we?")
rows = [("Strong", "The DEX is much more circular than English dictionaries; the cause is reciprocal (synonym) definition.", BLUE),
        ("Strong", "Swadesh and Graur words are concentrated in the Kernel (p < 10⁻²⁰⁰).", BLUE),
        ("Good", f"MinSet size: proven bounds {LB:,}–{UB:,}, a gap of {100*(UB-LB)/UB:.1f}%; not computed exactly.", AQUA),
        ("Good", "Graur’s most certain words are the most central; two marks out of 1,419 uncertain in our transcription.", AQUA),
        ("Indicative", "Etymological contrast: part of the vocabulary is unclassified.", ORANGE),
        ("Scope", "One dictionary, one edition, first senses only.", RGBColor(0x8a, 0x89, 0x84))]
for i, (tag, t, col) in enumerate(rows):
    y = Inches(1.9 + i * 0.78)
    b = box(s, Inches(0.7), y, Inches(1.8), Inches(0.6), col)
    tf = b.text_frame; tf.text = tag; r = tf.paragraphs[0].runs[0]
    r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = SURF; r.font.name = FONT
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, Inches(2.75), y + Inches(0.05), Inches(10.2), Inches(0.6), t, size=18)
notes(s, "Cât de siguri suntem: foarte siguri pe circularitate și pe concentrarea listelor de experți în nucleu; bine pe mărimea setului minim și pe poziția cuvintelor celor mai sigure ale lui Graur; "
         "orientativ pe contrastul etimologic. Limita principală: un singur dicționar.")

# 13 ── concluzie
s = base("Conclusion", "Summary")
text(s, Inches(0.7), Inches(1.9), Inches(12.0), Inches(4.0),
     [f"A third of the DEX forms a Kernel; about one word in ten must be known in advance to read the whole dictionary.",
      "About ten times the English proportion, mainly because the DEX defines words by synonyms.",
      "The Kernel contains almost all Swadesh and Graur words; Graur’s most certain words are the most central.",
      "A second, modern core of learned words does the defining.",
      ({"t": "We release the graph, its layers, the mutual pairs and the code.", "bold": True, "color": BLUE},)],
     size=22)
text(s, Inches(0.7), Inches(6.1), Inches(12), Inches(0.5),
     "vlad.belciug@365.univ-ovidius.ro · epelican@univ-ovidius.ro · Thanks to the dexonline volunteers.", size=15, color=INK2)
notes(s, "Concluzia: o treime din DEX formează nucleul; aproximativ un cuvânt din zece trebuie știut dinainte. De zece ori mai mult decât în engleză, din cauza definițiilor prin sinonim. Mulțumesc!")

prs.save("HiddenCore_ConsILR2026.pptx")
print("scris HiddenCore_ConsILR2026.pptx,", len(prs.slides), "slide-uri")
